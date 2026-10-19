"""Create the submission DOCX from the versioned specification."""
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2.2)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
for style in doc.styles:
    for border in list(style.element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr')):
        border.getparent().remove(border)
for line in (root / 'docs' / 'Техническое задание.md').read_text(encoding='utf-8').splitlines():
    if line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], style='Heading 1')
    elif line.strip():
        doc.add_paragraph(line)
doc.core_properties.title = 'Техническое задание на КП студию'
doc.core_properties.author = ''
doc.save(root / 'docs' / 'Техническое задание.docx')
print('DOCX создан')
